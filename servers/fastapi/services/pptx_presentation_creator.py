import os
from typing import Dict, List, Optional
from lxml import etree
from services.html_to_text_runs_service import (
    parse_html_text_to_text_runs as parse_inline_html_to_runs,
)

from pptx import Presentation
from pptx.shapes.autoshape import Shape
from pptx.slide import Slide
from pptx.text.text import _Paragraph, TextFrame, Font, _Run
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.enum.text import PP_ALIGN
from lxml.etree import fromstring, tostring
from PIL import Image
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

from pptx.util import Pt
from pptx.dml.color import RGBColor

from models.pptx_models import (
    PptxAutoShapeBoxModel,
    PptxBoxShapeEnum,
    PptxConnectorModel,
    PptxFillModel,
    PptxFontModel,
    PptxParagraphModel,
    PptxPictureBoxModel,
    PptxPositionModel,
    PptxPresentationModel,
    PptxShadowModel,
    PptxSlideModel,
    PptxSpacingModel,
    PptxStrokeModel,
    PptxStructureModel,
    PptxTableModel,
    PptxTextBoxModel,
    PptxTextRunModel,
)
from utils.download_helpers import download_files
from utils.image_utils import (
    clip_image,
    create_circle_image,
    fit_image,
    invert_image,
    round_image_corners,
    set_image_opacity,
)
import uuid

BLANK_SLIDE_LAYOUT = 6
CONTENT_PLACEHOLDER_IDX = 1


class PptxPresentationCreator:
    """
    Create a PowerPoint (.pptx) presentation from a structured presentation model.

    This class translates PptxPresentationModel and its nested slide and shape
    models into a concrete PowerPoint file using python-pptx. It supports both
    template-based placeholder rendering and absolute-positioned free shapes.
    """

    def __init__(
        self,
        ppt_model: PptxPresentationModel,
        temp_dir: str,
        template_path: Optional[str] = None,
    ):
        """
        Initialize the presentation creator.

        Args:
            ppt_model: Root presentation model containing slides and global shapes.
            temp_dir: Temporary directory for downloaded or transformed assets.
            template_path: Optional path to a PPTX template file.
        """
        self._temp_dir = temp_dir
        self._ppt_model = ppt_model
        self._slide_models = ppt_model.slides

        if template_path and os.path.exists(template_path):
            self._ppt = Presentation(template_path)
        else:
            self._ppt = Presentation()
            self._ppt.slide_width = Pt(1280)
            self._ppt.slide_height = Pt(720)

    # ==================== Main Flow ====================

    async def create_ppt(self):
        """
        Build the presentation content in memory.

        This method resolves all remote assets and iterates through slide models
        to create and populate slides. The resulting presentation must be saved
        explicitly using the `save` method.
        """
        await self.fetch_network_assets()

        for slide_model in self._slide_models:
            if self._ppt_model.shapes:
                slide_model.shapes.extend(self._ppt_model.shapes)

            self._add_slide(slide_model)

    def _add_slide(self, slide_model: PptxSlideModel):
        """
        Create and populate a single slide.

        Args:
            slide_model: Slide model defining layout, background, notes, and shapes.
        """
        layout_index = getattr(slide_model, "layout_index", BLANK_SLIDE_LAYOUT)
        slide = self._ppt.slides.add_slide(self._ppt.slide_layouts[layout_index])

        if slide_model.background:
            self._apply_fill(slide.background, slide_model.background)

        if slide_model.note:
            slide.notes_slide.notes_text_frame.text = slide_model.note

        placeholder_textboxes: List[PptxTextBoxModel] = []
        free_shapes = []

        for shape in slide_model.shapes:
            if isinstance(shape, PptxTextBoxModel) and shape.structure:
                placeholder_textboxes.append(shape)
            else:
                free_shapes.append(shape)

        if placeholder_textboxes:
            placeholder = self._get_placeholder(slide, CONTENT_PLACEHOLDER_IDX)
            if placeholder:
                self._fill_placeholder(placeholder.text_frame, placeholder_textboxes)
            else:
                free_shapes.extend(placeholder_textboxes)

        for shape in free_shapes:
            self._add_shape(slide, shape)

    # ==================== Placeholder ====================

    def _get_placeholder(self, slide: Slide, idx: int):
        """
        Retrieve a placeholder by index from a slide.

        Args:
            slide: Target slide.
            idx: Placeholder index.

        Returns:
            The placeholder shape if found, otherwise None.
        """
        for shape in slide.placeholders:
            if shape.placeholder_format.idx == idx:
                return shape
        return None

    def _fill_placeholder(
        self, text_frame: TextFrame, textboxes: List[PptxTextBoxModel]
    ):
        """
        Populate a placeholder text frame while preserving template bullet styles.

        Args:
            text_frame: Placeholder text frame to populate.
            textboxes: Text box models mapped to this placeholder.
        """
        first_para = True

        for textbox in textboxes:
            level = textbox.structure.level
            is_list = textbox.structure.isList

            for para_model in textbox.paragraphs:
                if first_para:
                    para = text_frame.paragraphs[0]
                    first_para = False
                else:
                    para = text_frame.add_paragraph()

                para.level = level
                para.alignment = PP_ALIGN.LEFT

                if not is_list:
                    self._disable_bullet(para)

                self._populate_paragraph(para, para_model)

    def _disable_bullet(self, paragraph: _Paragraph):
        """
        Disable bullet formatting for a paragraph.

        Args:
            paragraph: Target paragraph.
        """
        pPr = paragraph._p.get_or_add_pPr()

        for tag in ["a:buFont", "a:buChar", "a:buAutoNum"]:
            elem = pPr.find(qn(tag))
            if elem is not None:
                pPr.remove(elem)

        if pPr.find(qn("a:buNone")) is None:
            pPr.append(OxmlElement("a:buNone"))

    # ==================== Free-position Shapes ====================

    def _add_shape(self, slide: Slide, shape):
        """
        Dispatch shape creation based on its model type.

        Args:
            slide: Target slide.
            shape: Shape model instance.
        """
        if isinstance(shape, PptxPictureBoxModel):
            self._add_picture(slide, shape)
        elif isinstance(shape, PptxAutoShapeBoxModel):
            self._add_autoshape(slide, shape)
        elif isinstance(shape, PptxTextBoxModel):
            self._add_textbox(slide, shape)
        elif isinstance(shape, PptxConnectorModel):
            self._add_connector(slide, shape)
        elif isinstance(shape, PptxTableModel):
            self._add_table(slide, shape)

    def _add_textbox(self, slide: Slide, model: PptxTextBoxModel):
        """
        Add a textbox shape to a slide.

        Args:
            slide: Target slide.
            model: Text box model.
        """
        shape = slide.shapes.add_textbox(*model.position.to_pt_list())
        shape.width += Pt(2)

        tf = shape.text_frame
        tf.word_wrap = model.text_wrap

        self._apply_fill(shape, model.fill)
        self._apply_margin(tf, model.margin)
        self._add_paragraphs(tf, model.paragraphs)

    def _add_autoshape(self, slide: Slide, model: PptxAutoShapeBoxModel):
        """
        Add an auto shape to a slide.

        Args:
            slide: Target slide.
            model: Auto shape model.
        """
        position = self._get_margined_position(model.position, model.margin)
        shape = slide.shapes.add_shape(model.type, *position.to_pt_list())

        tf = shape.text_frame
        tf.word_wrap = model.text_wrap

        self._apply_fill(shape, model.fill)
        self._apply_margin(tf, model.margin)
        self._apply_stroke(shape, model.stroke)
        self._apply_shadow(shape, model.shadow)
        self._apply_border_radius(shape, model.border_radius)

        if model.paragraphs:
            self._add_paragraphs(tf, model.paragraphs)

    def _add_picture(self, slide: Slide, model: PptxPictureBoxModel):
        """
        Add a picture shape to a slide, applying optional image transformations.

        Args:
            slide: Target slide.
            model: Picture box model.

        """
        image_path = model.picture.path

        final_position = self._get_margined_position(model.position, model.margin)

        needs_processing = (
            model.clip
            or model.border_radius
            or model.invert
            or model.opacity
            or (model.object_fit and model.object_fit.fit)  # 只有 fit 有值才算
            or model.shape
        )

        if needs_processing:
            try:
                image = Image.open(image_path).convert("RGBA")
            except Exception:
                return

            if model.object_fit and model.object_fit.fit:
                # object_fit.fit 有明確值時才使用 fit_image
                image = fit_image(
                    image,
                    final_position.width,
                    final_position.height,
                    model.object_fit,
                )
            elif model.clip:
                image = clip_image(
                    image,
                    final_position.width,
                    final_position.height,
                )

            if model.border_radius:
                image = round_image_corners(image, model.border_radius)
            if model.shape == PptxBoxShapeEnum.CIRCLE:
                image = create_circle_image(image)
            if model.invert:
                image = invert_image(image)
            if model.opacity:
                image = set_image_opacity(image, model.opacity)

            image_path = os.path.join(self._temp_dir, f"{uuid.uuid4()}.png")
            image.save(image_path)

        slide.shapes.add_picture(image_path, *final_position.to_pt_list())

    def _add_connector(self, slide: Slide, model: PptxConnectorModel):
        """
        Add a connector (line) shape to a slide.

        Args:
            slide: Target slide.
            model: Connector model.
        """
        if model.thickness == 0:
            return
        shape = slide.shapes.add_connector(
            model.type, *model.position.to_pt_xyxy()
        )
        shape.line.width = Pt(model.thickness)
        shape.line.color.rgb = RGBColor.from_string(model.color)
        self._set_fill_opacity(shape, model.opacity)

    def _add_table(self, slide: Slide, model: PptxTableModel):
        """
        Add a table shape to a slide.

        Args:
            slide: Target slide.
            model: Table model containing rows of cells.
        """
        if not model.rows:
            return

        num_rows = len(model.rows)
        num_cols = len(model.rows[0]) if model.rows else 0

        if num_rows == 0 or num_cols == 0:
            return

        # Create table with position
        table_shape = slide.shapes.add_table(
            num_rows,
            num_cols,
            Pt(model.position.left),
            Pt(model.position.top),
            Pt(model.position.width),
            Pt(model.position.height),
        )
        table = table_shape.table

        # Default fonts
        default_font = model.font or PptxFontModel(name="Inter", size=12, color="000000")
        header_font = model.header_font or PptxFontModel(
            name="Inter", size=12, color="FFFFFF", font_weight=700
        )

        # Populate cells
        for row_idx, row in enumerate(model.rows):
            is_header = model.header_row and row_idx == 0

            for col_idx, cell_model in enumerate(row):
                if col_idx >= num_cols:
                    break

                cell = table.cell(row_idx, col_idx)
                cell.text = cell_model.text

                # Apply cell fill
                if is_header and model.header_fill:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(model.header_fill.color)
                elif not is_header and model.cell_fill:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(model.cell_fill.color)

                # Apply font to paragraphs
                font_model = header_font if is_header else default_font
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.name = font_model.name
                    paragraph.font.size = Pt(font_model.size)
                    paragraph.font.color.rgb = RGBColor.from_string(font_model.color)
                    paragraph.font.bold = font_model.font_weight >= 600 if font_model.font_weight else False

    # ==================== Paragraph / TextRun ====================

    def _add_paragraphs(
        self, text_frame: TextFrame, models: List[PptxParagraphModel]
    ):
        """
        Add paragraphs to a text frame.

        Args:
            text_frame: Target text frame.
            models: Paragraph models.
        """
        for i, model in enumerate(models):
            para = (
                text_frame.paragraphs[0]
                if i == 0
                else text_frame.add_paragraph()
            )
            self._populate_paragraph(para, model)

    def _populate_paragraph(
        self, para: _Paragraph, model: PptxParagraphModel
    ):
        """
        Populate a paragraph with text runs and styles.

        Args:
            para: Target paragraph.
            model: Paragraph model.
        """
        if model.spacing:
            para.space_before = Pt(model.spacing.top)
            para.space_after = Pt(model.spacing.bottom)

        if model.line_height:
            para.line_spacing = model.line_height

        if model.alignment:
            para.alignment = model.alignment

        if model.font:
            self._apply_font(para.font, model.font)

        text_runs = []
        if model.text:
            text_runs = parse_inline_html_to_runs(model.text, model.font)
        elif model.text_runs:
            text_runs = model.text_runs

        for run_model in text_runs:
            run = para.add_run()
            run.text = run_model.text
            if run_model.font:
                self._apply_font(run.font, run_model.font)

    # ==================== Style Helpers ====================

    def _apply_font(self, font: Font, model: PptxFontModel):
        """
        Apply font styling to a text element.

        Args:
            font: Target font object.
            model: Font model.
        """
        font.name = model.name
        font.size = Pt(model.size)
        font.color.rgb = RGBColor.from_string(model.color)
        font.italic = model.italic
        font.bold = model.font_weight >= 600

        if model.underline is not None:
            font.underline = bool(model.underline)
        if model.strike is not None:
            rPr = font._element
            rPr.set(
                "strike", "sngStrike" if model.strike else "noStrike"
            )

    def _apply_fill(self, shape: Shape, fill: Optional[PptxFillModel]):
        """
        Apply fill color and opacity to a shape.

        Args:
            shape: Target shape.
            fill: Fill model.
        """
        if not fill:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(
                fill.color
            )
            self._set_fill_opacity(shape.fill, fill.opacity)

    def _apply_stroke(
        self, shape: Shape, stroke: Optional[PptxStrokeModel]
    ):
        """
        Apply stroke (outline) styling to a shape.

        Args:
            shape: Target shape.
            stroke: Stroke model.
        """
        if not stroke or stroke.thickness == 0:
            shape.line.fill.background()
        else:
            shape.line.fill.solid()
            shape.line.fill.fore_color.rgb = RGBColor.from_string(
                stroke.color
            )
            shape.line.width = Pt(stroke.thickness)
            self._set_fill_opacity(shape.line.fill, stroke.opacity)

    def _apply_shadow(
        self, shape: Shape, shadow: Optional[PptxShadowModel]
    ):
        """
        Apply shadow effect to a shape.

        Args:
            shape: Target shape.
            shadow: Shadow model.
        """
        sp_pr = shape._element.xpath("p:spPr")[0]
        nsmap = sp_pr.nsmap

        effect_list = sp_pr.find("a:effectLst", namespaces=nsmap)
        if effect_list is not None:
            for tag in ["a:outerShdw", "a:innerShdw", "a:prstShdw"]:
                old = effect_list.find(tag, namespaces=nsmap)
                if old is not None:
                    effect_list.remove(old)
        else:
            effect_list = etree.SubElement(
                sp_pr,
                f"{{{nsmap['a']}}}effectLst",
                nsmap=nsmap,
            )

        if shadow is None:
            outer = etree.SubElement(
                effect_list,
                f"{{{nsmap['a']}}}outerShdw",
                {"blurRad": "0", "dist": "0", "dir": "0"},
                nsmap=nsmap,
            )
            color = etree.SubElement(
                outer,
                f"{{{nsmap['a']}}}srgbClr",
                {"val": "000000"},
                nsmap=nsmap,
            )
            etree.SubElement(
                color,
                f"{{{nsmap['a']}}}alpha",
                {"val": "0"},
                nsmap=nsmap,
            )
        else:
            angle = (
                int(round((shadow.angle % 360) * 60000))
                if shadow.angle
                else 0
            )
            outer = etree.SubElement(
                effect_list,
                f"{{{nsmap['a']}}}outerShdw",
                {
                    "blurRad": f"{Pt(shadow.radius)}",
                    "dir": f"{angle}",
                    "dist": f"{Pt(shadow.offset)}",
                    "rotWithShape": "0",
                },
                nsmap=nsmap,
            )
            color = etree.SubElement(
                outer,
                f"{{{nsmap['a']}}}srgbClr",
                {"val": shadow.color},
                nsmap=nsmap,
            )
            etree.SubElement(
                color,
                f"{{{nsmap['a']}}}alpha",
                {"val": f"{int(shadow.opacity * 100000)}"},
                nsmap=nsmap,
            )

    def _apply_border_radius(
        self, shape: Shape, radius: Optional[int]
    ):
        """
        Apply rounded corners to a shape if supported.

        Args:
            shape: Target shape.
            radius: Border radius value.
        """
        if not radius:
            return
        try:
            shape.adjustments[0] = Pt(radius) / min(
                shape.width, shape.height
            )
        except Exception:
            pass

    def _apply_margin(
        self, tf: TextFrame, margin: Optional[PptxSpacingModel]
    ):
        """
        Apply text margins to a text frame.

        Args:
            tf: Target text frame.
            margin: Spacing model.
        """
        tf.margin_left = Pt(margin.left if margin else 0)
        tf.margin_right = Pt(margin.right if margin else 0)
        tf.margin_top = Pt(margin.top if margin else 0)
        tf.margin_bottom = Pt(margin.bottom if margin else 0)

    def _set_fill_opacity(self, fill, opacity):
        """
        Set opacity on a fill element.

        Args:
            fill: Fill object.
            opacity: Opacity value between 0.0 and 1.0.
        """
        if opacity is None or opacity >= 1.0:
            return
        try:
            sF = fill._xPr.solidFill.get_or_change_to_srgbClr()
            elem = OxmlElement("a:alpha")
            elem.set("val", str(int(opacity * 100000)))
            sF.append(elem)
        except Exception:
            pass

    def _get_margined_position(
        self,
        pos: PptxPositionModel,
        margin: Optional[PptxSpacingModel],
    ) -> PptxPositionModel:
        """
        Compute a position adjusted by margins.

        Args:
            pos: Original position.
            margin: Margin values.

        Returns:
            A new position model with margins applied.
        """
        if not margin:
            return pos
        return PptxPositionModel(
            left=pos.left + margin.left,
            top=pos.top + margin.top,
            width=max(pos.width - margin.left - margin.right, 0),
            height=max(pos.height - margin.top - margin.bottom, 0),
        )

    # ==================== Network Assets ====================

    async def fetch_network_assets(self):
        """
        Download all remote image assets referenced by the presentation model.

        Image URLs are replaced with local file paths under `temp_dir`.
        """
        image_urls = []
        models_with_network_asset: List[PptxPictureBoxModel] = []

        all_shapes = list(self._ppt_model.shapes or [])
        for slide in self._slide_models:
            all_shapes.extend(slide.shapes)

        for shape in all_shapes:
            if isinstance(shape, PptxPictureBoxModel):
                path = shape.picture.path
                if path.startswith("http"):
                    if "app_data" in path:
                        shape.picture.path = os.path.join(
                            "/app_data",
                            path.split("app_data/")[1],
                        )
                        shape.picture.is_network = False
                    else:
                        image_urls.append(path)
                        models_with_network_asset.append(shape)

        if image_urls:
            paths = await download_files(image_urls, self._temp_dir)
            for shape, path in zip(models_with_network_asset, paths):
                if path:
                    shape.picture.path = path
                    shape.picture.is_network = False

    # ==================== Save ====================

    def save(self, path: str):
        """
        Save the generated presentation to disk.

        Args:
            path: Output file path.
        """
        self._ppt.save(path)
